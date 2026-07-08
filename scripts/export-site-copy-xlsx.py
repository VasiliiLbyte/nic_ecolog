#!/usr/bin/env python3
"""Convert site copy review Word doc to Excel for Google Sheets iterations."""

from __future__ import annotations

import argparse
import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "exports"
DEFAULT_DOCX = Path.home() / "Downloads" / "НИЦ Эколог — контент сайта для вычитки (1).docx"
FALLBACK_DOCX = OUT_DIR / "НИЦ Эколог — контент сайта для вычитки.docx"

ITERATION_COLUMNS = [
    "Вычитка 1",
    "Вычитка 2",
    "Вычитка 3",
]

HEADERS = [
    "ID",
    "Страница",
    "Раздел",
    "Тип",
    "Текст на сайте",
    "Примечание",
    *ITERATION_COLUMNS,
    "Статус",
]

STATUS_OPTIONS = "черновик,на проверке,принято,внедрено,отложено"


def iter_block_items(parent):
    """Yield paragraphs and tables in document order."""
    if hasattr(parent, "element"):
        parent_el = parent.element.body if hasattr(parent, "element") and parent.element.tag.endswith("document") else parent.element
    else:
        parent_el = parent

    for child in parent_el.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            yield Paragraph(child, parent)
        elif tag == "tbl":
            yield Table(child, parent)


def heading_level(paragraph: Paragraph) -> int | None:
    name = paragraph.style.name if paragraph.style else ""
    if name == "Title":
        return 0
    m = re.match(r"Heading\s+(\d+)", name)
    return int(m.group(1)) if m else None


def split_site_text(raw: str) -> tuple[str, str]:
    """Separate visible text and trailing note like (#anchor) or (tel:...)."""
    raw = raw.strip()
    note = ""
    m = re.search(r"\n\(([^)]+)\)\s*$", raw)
    if m:
        note = m.group(1).strip()
        raw = raw[: m.start()].strip()
    return raw, note


def parse_docx(path: Path) -> list[dict]:
    doc = Document(path)
    rows: list[dict] = []
    current_page = ""
    current_section = ""
    row_num = 0

    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            level = heading_level(block)
            text = block.text.strip()
            if level == 1 and text:
                current_page = text
                current_section = ""
            elif level == 2 and text:
                current_section = text
            continue

        if not isinstance(block, Table):
            continue

        table_rows = block.rows
        if not table_rows:
            continue

        header = [c.text.strip() for c in table_rows[0].cells]
        if len(header) < 2 or header[0] != "Тип":
            continue

        for tr in table_rows[1:]:
            cells = [c.text.strip() for c in tr.cells]
            while len(cells) < 3:
                cells.append("")

            kind, site_raw, edit_raw = cells[0], cells[1], cells[2]
            if not kind and not site_raw and not edit_raw:
                continue

            site_text, note = split_site_text(site_raw)
            row_num += 1
            rows.append(
                {
                    "id": f"R{row_num:04d}",
                    "page": current_page,
                    "section": current_section,
                    "kind": kind,
                    "site_text": site_text,
                    "note": note,
                    "iterations": [edit_raw] + [""] * (len(ITERATION_COLUMNS) - 1),
                }
            )

    return rows


def style_sheet(ws, row_count: int) -> None:
    header_fill = PatternFill("solid", fgColor="0B5C60")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    body_font = Font(name="Calibri", size=11)
    wrap = Alignment(wrap_text=True, vertical="top")

    widths = {
        "A": 8,
        "B": 22,
        "C": 34,
        "D": 16,
        "E": 52,
        "F": 18,
        "G": 52,
        "H": 52,
        "I": 52,
        "J": 14,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    for col_idx in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx in range(2, row_count + 2):
        for col_idx in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = body_font
            cell.alignment = wrap

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{row_count + 1}"

    status_col = get_column_letter(HEADERS.index("Статус") + 1)
    dv = DataValidation(type="list", formula1=f'"{STATUS_OPTIONS}"', allow_blank=True)
    dv.error = "Выберите значение из списка"
    dv.errorTitle = "Статус"
    ws.add_data_validation(dv)
    dv.add(f"{status_col}2:{status_col}{row_count + 1}")


def build_xlsx(docx_path: Path, out_path: Path) -> Path:
    rows = parse_docx(docx_path)
    if not rows:
        raise SystemExit(f"Не удалось извлечь строки из {docx_path}")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Контент"

    ws.append(HEADERS)
    for row in rows:
        ws.append(
            [
                row["id"],
                row["page"],
                row["section"],
                row["kind"],
                row["site_text"],
                row["note"],
                *row["iterations"],
                "",
            ]
        )

    style_sheet(ws, len(rows))
    wb.save(out_path)
    return out_path


def resolve_docx(path: str | None) -> Path:
    if path:
        p = Path(path).expanduser()
        if not p.exists():
            raise SystemExit(f"Файл не найден: {p}")
        return p
    if DEFAULT_DOCX.exists():
        return DEFAULT_DOCX
    if FALLBACK_DOCX.exists():
        return FALLBACK_DOCX
    raise SystemExit(
        "Укажите путь к .docx: --input /path/to/file.docx\n"
        f"Ожидался также: {DEFAULT_DOCX}"
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Конвертация Word-вычитки в Excel для Google Sheets")
    parser.add_argument("--input", "-i", help="Путь к .docx (по умолчанию — файл из Downloads)")
    parser.add_argument(
        "--output",
        "-o",
        default=str(OUT_DIR / "НИЦ Эколог — контент сайта.xlsx"),
        help="Путь к выходному .xlsx",
    )
    args = parser.parse_args()

    docx_path = resolve_docx(args.input)
    out_path = Path(args.output).expanduser()
    result = build_xlsx(docx_path, out_path)
    parsed = parse_docx(docx_path)
    edits = sum(1 for r in parsed if r["iterations"][0])
    print(f"Создан файл: {result}")
    print(f"Строк: {len(parsed)}, с правками в «Вычитка 1»: {edits}")


if __name__ == "__main__":
    main()
