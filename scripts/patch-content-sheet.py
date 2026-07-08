#!/usr/bin/env python3
"""Patch content spreadsheet rows (one-off / maintenance)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "exports" / "НИЦ Эколог — контент сайта.xlsx"

STEP2 = (
    "Расчёт выбросов загрязняющих веществ и/или проведение "
    "инструментальных замеров на площадке при необходимости"
)
STEP3 = (
    "Проведение рассеивания загрязняющих веществ в атмосферном воздухе "
    "с целью оценки воздействия деятельности предприятия на ближайшие нормируемые объекты"
)


def col_index(ws, header: str) -> int:
    for idx, cell in enumerate(ws[1], start=1):
        if cell.value == header or (header in str(cell.value) and header == "Вычитка 1"):
            return idx
    raise KeyError(header)


def find_row(ws, row_id: str) -> int:
    id_col = col_index(ws, "ID")
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, id_col).value == row_id:
            return r
    raise KeyError(row_id)


def patch_ndv_steps() -> None:
    wb = load_workbook(XLSX)
    if "Инструкция" in wb.sheetnames:
        del wb["Инструкция"]

    ws = wb["Контент"]
    # Align header with Google Sheets
    g_col = col_index(ws, "Вычитка 1")
    ws.cell(1, g_col).value = "Вычитка 1 01.07.2026"

    type_col = col_index(ws, "Тип")
    text_col = col_index(ws, "Текст на сайте")
    iter_col = g_col

    r272 = find_row(ws, "R0272")
    template = [ws.cell(r272, c).value for c in range(1, ws.max_column + 1)]

    # Шаг 2 — только первый абзац
    ws.cell(r272, type_col).value = "Шаг 2"
    ws.cell(r272, text_col).value = STEP2
    ws.cell(r272, iter_col).value = STEP2

    # Вставка Шаг 3 между R0272 и R0273
    insert_at = r272 + 1
    ws.insert_rows(insert_at)
    new_row = template.copy()
    new_row[0] = "R0272b"  # ID
    new_row[type_col - 1] = "Шаг 3"
    new_row[text_col - 1] = STEP3
    new_row[iter_col - 1] = STEP3
    for c, val in enumerate(new_row, start=1):
        ws.cell(insert_at, c).value = val

    # Перенумерация следующих шагов
    renames = {
        "R0273": ("Шаг 4", "Формирование проекта НДВ", "Формирование проекта НДВ"),
        "R0274": ("Шаг 5", None, None),
        "R0275": ("Шаг 6", None, None),
    }
    for rid, (step_type, site_text, iter_text) in renames.items():
        r = find_row(ws, rid)
        ws.cell(r, type_col).value = step_type
        if site_text is not None:
            ws.cell(r, text_col).value = site_text
        if iter_text is not None:
            ws.cell(r, iter_col).value = iter_text

    wb.save(XLSX)
    print(f"Обновлено: {XLSX}")


if __name__ == "__main__":
    patch_ndv_steps()
